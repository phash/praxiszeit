"""Release-Review 1.20.1: ODS-Spalte 10 „Bemerkung" pariert mit XLSX.

XLSX baut die Bemerkung eines Tages aus den Zeiteinträgen (§10-Ausnahmegrund
nur an Sonn-/Feiertagen, dann die Notiz) und überschreibt sie im
Abwesenheitszweig NUR, wenn die Abwesenheiten eine Notiz tragen. ODS schrieb im
Abwesenheitszweig immer die (oft leere) Abwesenheitsnotiz und außerhalb des
Beschäftigungsfensters immer "" — die Notiz des Zeiteintrags fehlte an einem
Misch-Tag (Zeiteintrag „Arzttermin" + halber Tag Urlaub ohne Notiz).
Gilt für das Monatsblatt und das Jahres-Mitarbeiterblatt.
"""
from datetime import date, datetime, time
from io import BytesIO

import pytest
from openpyxl import load_workbook

from app.models import Absence, AbsenceType, TimeEntry, User, UserRole
from app.services import export_service, ods_export_service
from tests.conftest import DEFAULT_TENANT_ID

YEAR, MONTH = 2026, 9
BEFORE = date(2026, 9, 1)        # Di — vor dem Eintritt, Eintrag mit Notiz
FIRST_WORK_DAY = date(2026, 9, 2)
MIXED = date(2026, 9, 3)         # Do — Eintrag „Arzttermin" + Halbtag Urlaub ohne Notiz
MIXED_NOTE = date(2026, 9, 4)    # Fr — Eintrag mit Notiz + Urlaub MIT Notiz
SATURDAY = date(2026, 9, 5)      # Sa — §10-Grund an einem Samstag (kein Feiertag)
FORMULA = date(2026, 9, 8)       # Di — Formel-Notiz + Halbtag ohne Notiz


@pytest.fixture
def emp(db, default_tenant):
    from app.services import auth_service
    u = User(
        username="remark", email="remark@example.com",
        password_hash=auth_service.hash_password("test123"),
        first_name="Bem", last_name="Erkung",
        role=UserRole.EMPLOYEE, weekly_hours=40.0, vacation_days=30,
        work_days_per_week=5, is_active=True, tenant_id=DEFAULT_TENANT_ID,
        first_work_day=FIRST_WORK_DAY,
    )
    db.add(u)
    db.commit()
    db.refresh(u)

    def entry(d, note=None, sunday_reason=None, start=time(8, 0), end=time(12, 0)):
        db.add(TimeEntry(user_id=u.id, tenant_id=DEFAULT_TENANT_ID, date=d,
                         start_time=start, end_time=end, break_minutes=0,
                         note=note, sunday_exception_reason=sunday_reason))

    def absence(d, note=None):
        db.add(Absence(user_id=u.id, tenant_id=DEFAULT_TENANT_ID, date=d,
                       type=AbsenceType.VACATION, hours=4.0, half_day=True, note=note))

    entry(BEFORE, note="Probetag")
    entry(MIXED, note="Arzttermin")
    absence(MIXED)
    entry(MIXED_NOTE, note="Eintragsnotiz")
    absence(MIXED_NOTE, note="Urlaubsnotiz")
    entry(SATURDAY, note="Inventur", sunday_reason="Notdienst")
    entry(FORMULA, note="=1+1")
    absence(FORMULA)
    db.commit()
    return u


def _xlsx_remarks(bio: BytesIO, sheet_name):
    bio.seek(0)
    sheet = load_workbook(bio)[sheet_name]
    out = {}
    for r in sheet.iter_rows(min_col=1, max_col=10, values_only=True):
        d = r[0].date() if isinstance(r[0], datetime) else r[0]
        if isinstance(d, date):
            out[d] = r[9] or ""
    return out


def _ods_remarks(bio: BytesIO, sheet_name):
    from odf.opendocument import load as odf_load
    from odf.table import Table, TableCell, TableRow
    from odf.teletype import extractText
    bio.seek(0)
    doc = odf_load(bio)
    table = [t for t in doc.spreadsheet.getElementsByType(Table)
             if t.getAttribute("name") == sheet_name][0]
    out = {}
    for tr in table.getElementsByType(TableRow):
        cells = [extractText(c) for c in tr.getElementsByType(TableCell)]
        if not cells:
            continue
        try:
            d = datetime.strptime(cells[0], "%d.%m.%Y").date()
        except ValueError:
            continue
        out[d] = cells[9] if len(cells) > 9 else ""
    return out


DAYS = [BEFORE, MIXED, MIXED_NOTE, SATURDAY, FORMULA]


def test_monthly_remark_matches_xlsx(db, emp):
    xlsx = _xlsx_remarks(
        export_service.generate_monthly_report(db, YEAR, MONTH, tenant_id=DEFAULT_TENANT_ID),
        f"{emp.last_name} {emp.first_name}"[:31])
    ods = _ods_remarks(
        ods_export_service.generate_monthly_report(db, YEAR, MONTH, tenant_id=DEFAULT_TENANT_ID),
        f"{emp.last_name} {emp.first_name}"[:31])
    # Kontrolle des Szenarios (XLSX ist die Referenz)
    assert xlsx[MIXED] == "Arzttermin"
    assert xlsx[BEFORE] == "Probetag"
    assert xlsx[MIXED_NOTE] == "Urlaubsnotiz"
    assert {d: ods[d] for d in DAYS} == {d: xlsx[d] for d in DAYS}


def test_yearly_employee_remark_matches_xlsx(db, emp):
    xlsx = _xlsx_remarks(
        export_service.generate_yearly_report(db, YEAR, tenant_id=DEFAULT_TENANT_ID),
        emp.last_name[:20])
    ods = _ods_remarks(
        ods_export_service.generate_yearly_report(db, YEAR, tenant_id=DEFAULT_TENANT_ID),
        f"{emp.last_name} {emp.first_name}"[:31])
    assert xlsx[MIXED] == "Arzttermin"
    assert xlsx[BEFORE] == "Probetag"
    assert {d: ods[d] for d in DAYS} == {d: xlsx[d] for d in DAYS}
