"""#495 — Der Schichtplan-Aushang zeigte on-prem „Default" statt des Praxisnamens.

Der On-Prem-Bootstrap (``main.py``) und schon Migration ``027`` legen den
einzigen Mandanten mit ``name="Default"`` an. Der beim Installieren abgefragte
Praxisname landet dagegen in ``settings.PRACTICE_NAME`` (``[practice] name`` in
``config/praxiszeit.conf`` bzw. ``PRACTICE_NAME`` in der Docker-``.env``). Der
PDF-Aushang las ``Tenant.name`` — und on-prem gibt es keine Oberfläche, die
diesen Wert ändert. „Default" stand damit auf einem Ausdruck, der für alle
Mitarbeitenden sichtbar am Schwarzen Brett hängt.

Fix: ``practice_name_service.practice_display_name(tenant)`` ist die EINE
Antwort auf „wie heißt die Praxis auf einem Ausdruck?". On-prem gilt der
Bootstrap-Wert „Default" (zusammen mit dem Slug ``default``) als Platzhalter
und wird durch den konfigurierten Praxisnamen ersetzt — zur Laufzeit, ohne
Schreibzugriff auf die ``tenants``-Zeile (Slug bleibt Bootstrap-Anker, #435).
"""

from io import BytesIO
from pathlib import Path

from app.config import settings
from app.models import UserRole
from app.models.tenant import Tenant
from app.routers.shift_planning import export_plan_pdf
from tests.test_shift_plan_pdf import _body, _pdf_text, _plan, _user

DEMO_NAME = "Hausarztpraxis Dr. Muster"


# ─── Endpunkt: der Aushang ──────────────────────────────────────────────


def test_shift_plan_pdf_passes_configured_practice_name_on_prem(db, default_tenant, monkeypatch):
    """Der Endpunkt reicht dem (datenbankfreien) Renderer den Praxisnamen
    hinein. On-prem mit unverändertem Bootstrap-Mandanten muss das der
    konfigurierte Name sein, nicht „Default"."""
    from app.services import shift_plan_export_service

    monkeypatch.setattr(settings, "DEPLOYMENT_MODE", "onprem")
    monkeypatch.setattr(settings, "PRACTICE_NAME", DEMO_NAME)
    admin = _user(db, "pn_admin", role=UserRole.ADMIN)
    plan = _plan(db, admin, "Aushang")

    captured = {}

    def _fake(detail, *, practice_name, **kwargs):
        captured["practice_name"] = practice_name
        return BytesIO(b"%PDF-fake")

    monkeypatch.setattr(shift_plan_export_service, "generate_plan_pdf", _fake)

    export_plan_pdf(plan.id, db=db, current_user=admin)

    assert captured["practice_name"] == DEMO_NAME


def test_shift_plan_pdf_header_shows_practice_name_not_default(db, default_tenant, monkeypatch):
    """Ende-zu-Ende über den echten Renderer: genau die Kopfzeile aus dem
    Fehlerbericht („Default · Standort: …") darf nicht mehr entstehen."""
    monkeypatch.setattr(settings, "DEPLOYMENT_MODE", "onprem")
    monkeypatch.setattr(settings, "PRACTICE_NAME", DEMO_NAME)
    admin = _user(db, "pn_e2e_admin", role=UserRole.ADMIN)
    plan = _plan(db, admin, "Aushang Ende-zu-Ende")

    text = _pdf_text(BytesIO(_body(export_plan_pdf(plan.id, db=db, current_user=admin))))

    assert DEMO_NAME in text
    assert "Default" not in text


def test_shift_plan_pdf_keeps_a_real_tenant_name(db, monkeypatch):
    """SaaS: der Mandantenname IST der beim Signup eingegebene Praxisname —
    der dort gesetzte Wert gewinnt, ``PRACTICE_NAME`` (eine Einstellung des
    Betreibers, nicht der Praxis) darf ihn nicht überschreiben."""
    from app.services import shift_plan_export_service
    from tests.conftest import DEFAULT_TENANT_ID

    monkeypatch.setattr(settings, "DEPLOYMENT_MODE", "saas")
    monkeypatch.setattr(settings, "PRACTICE_NAME", "Betreiber-Platzhalter")
    db.add(Tenant(id=DEFAULT_TENANT_ID, name="Zahnarztpraxis Sonnenschein",
                  slug="zahnarztpraxis-sonnenschein", is_active=True, mode="multi"))
    db.commit()
    admin = _user(db, "pn_saas_admin", role=UserRole.ADMIN)
    plan = _plan(db, admin, "SaaS-Aushang")

    captured = {}

    def _fake(detail, *, practice_name, **kwargs):
        captured["practice_name"] = practice_name
        return BytesIO(b"%PDF-fake")

    monkeypatch.setattr(shift_plan_export_service, "generate_plan_pdf", _fake)

    export_plan_pdf(plan.id, db=db, current_user=admin)

    assert captured["practice_name"] == "Zahnarztpraxis Sonnenschein"


# ─── Zweite Fläche mit Tenant.name als Praxisname: der AVV-Entwurf ──────


def test_avv_names_the_configured_practice_on_prem(monkeypatch):
    """``avv_generator`` druckte ``tenant.company_name or tenant.name`` als
    „Verantwortlicher" — on-prem also ebenfalls „Default"."""
    from app.services.avv_generator import generate_avv_pdf

    monkeypatch.setattr(settings, "DEPLOYMENT_MODE", "onprem")
    monkeypatch.setattr(settings, "PRACTICE_NAME", DEMO_NAME)
    tenant = Tenant(name="Default", slug="default", is_active=True, mode="single")

    text = _pdf_text(BytesIO(generate_avv_pdf(tenant)))

    assert DEMO_NAME in text
    assert "Default" not in text


def test_avv_still_prefers_the_billing_company_name(monkeypatch):
    """Die Firmierung aus den Abrechnungsdaten bleibt vorrangig — der Fix
    ersetzt nur den Rückfallwert ``tenant.name``."""
    from app.services.avv_generator import generate_avv_pdf

    monkeypatch.setattr(settings, "DEPLOYMENT_MODE", "onprem")
    monkeypatch.setattr(settings, "PRACTICE_NAME", DEMO_NAME)
    tenant = Tenant(name="Default", slug="default", is_active=True, mode="single",
                    company_name="Muster MVZ GmbH")

    text = _pdf_text(BytesIO(generate_avv_pdf(tenant)))

    assert "Muster MVZ GmbH" in text


# ─── Der Helfer selbst ──────────────────────────────────────────────────


def _tenant(name, slug="default"):
    return Tenant(name=name, slug=slug, is_active=True, mode="single")


def test_helper_replaces_bootstrap_placeholder_on_prem(monkeypatch):
    from app.services.practice_name_service import practice_display_name

    monkeypatch.setattr(settings, "DEPLOYMENT_MODE", "onprem")
    monkeypatch.setattr(settings, "PRACTICE_NAME", f"  {DEMO_NAME}  ")

    assert practice_display_name(_tenant("Default")) == DEMO_NAME


def test_helper_keeps_any_other_stored_name_on_prem(monkeypatch):
    """Nur der Bootstrap-Wert ist ein Platzhalter. Ein bewusst gesetzter Name
    (oder der Anonymisierungs-Ersatz „[gelöscht]" nach Art. 17) wird nicht
    durch die Konfiguration überschrieben."""
    from app.services.practice_name_service import practice_display_name

    monkeypatch.setattr(settings, "DEPLOYMENT_MODE", "onprem")
    monkeypatch.setattr(settings, "PRACTICE_NAME", DEMO_NAME)

    assert practice_display_name(_tenant("Praxis am Markt")) == "Praxis am Markt"
    assert practice_display_name(_tenant("[gelöscht]")) == "[gelöscht]"


def test_helper_requires_the_default_slug(monkeypatch):
    """Platzhalter ist nur der Bootstrap-Mandant (Slug ``default``), nicht
    irgendein Mandant, der zufällig „Default" heißt."""
    from app.services.practice_name_service import practice_display_name

    monkeypatch.setattr(settings, "DEPLOYMENT_MODE", "onprem")
    monkeypatch.setattr(settings, "PRACTICE_NAME", DEMO_NAME)

    assert practice_display_name(_tenant("Default", slug="default-2")) == "Default"


def test_helper_does_not_touch_saas_tenants(monkeypatch):
    """SaaS: eine Praxis, die sich beim Signup „Default" nennt, heißt so."""
    from app.services.practice_name_service import practice_display_name

    monkeypatch.setattr(settings, "DEPLOYMENT_MODE", "saas")
    monkeypatch.setattr(settings, "PRACTICE_NAME", DEMO_NAME)

    assert practice_display_name(_tenant("Default")) == "Default"


def test_helper_omits_an_empty_configured_name(monkeypatch):
    """Ist kein Praxisname konfiguriert, lieber gar nichts in der Kopfzeile
    als „Default" — der Renderer lässt ``None`` aus."""
    from app.services.practice_name_service import practice_display_name

    monkeypatch.setattr(settings, "DEPLOYMENT_MODE", "onprem")
    monkeypatch.setattr(settings, "PRACTICE_NAME", "   ")

    assert practice_display_name(_tenant("Default")) is None
    assert practice_display_name(None) is None


def test_helper_does_not_write_to_the_tenant(monkeypatch):
    """Laufzeit-Ersatz, kein Datenumbau: die Zeile bleibt unangetastet, damit
    eine spätere Änderung von ``PRACTICE_NAME`` (Konfiguration + Neustart)
    ohne Migration greift und der Slug als Bootstrap-Anker stehen bleibt."""
    from app.services.practice_name_service import practice_display_name

    monkeypatch.setattr(settings, "DEPLOYMENT_MODE", "onprem")
    monkeypatch.setattr(settings, "PRACTICE_NAME", DEMO_NAME)
    tenant = _tenant("Default")

    practice_display_name(tenant)

    assert tenant.name == "Default"
    assert tenant.slug == "default"


def test_placeholder_matches_what_migration_027_and_bootstrap_write():
    """Bestandsinstallationen tragen den Wert aus Migration 027, Neuinstallationen
    den aus dem Bootstrap in ``main.py``. Beide müssen dem Platzhalter
    entsprechen, den der Helfer erkennt — sonst greift der Ersatz still nicht."""
    from app.services import practice_name_service

    backend = Path(__file__).resolve().parents[1]
    migration = (backend / "alembic" / "versions" / "2026_03_25_2000-027_add_multi_tenant.py").read_text(encoding="utf-8")
    main_py = (backend / "app" / "main.py").read_text(encoding="utf-8")

    assert practice_name_service.DEFAULT_TENANT_NAME == "Default"
    assert practice_name_service.DEFAULT_TENANT_SLUG == "default"
    assert "'Default', 'default'" in migration
    assert "name=practice_name_service.DEFAULT_TENANT_NAME" in main_py


# ─── Dritte Fläche: der klassische Jahresbericht (XLSX, Zelle A1) ───────
#
# ``_create_employee_classic_sheet`` schrieb ``settings.PRACTICE_NAME`` direkt
# in die Kopfzelle. On-prem ist das derselbe Wert, den der Helfer liefert —
# in SaaS aber die Einstellung des Betreibers (Vorgabe „Praxis") für JEDEN
# Mandanten, während Aushang und AVV desselben Mandanten dessen Signup-Namen
# zeigen. Ein Mandant, zwei Praxisnamen auf zwei Ausdrucken.


def _classic_a1(db):
    from app.services.export_service import generate_yearly_report_classic
    from openpyxl import load_workbook
    from tests.conftest import DEFAULT_TENANT_ID

    out = generate_yearly_report_classic(db, 2026, tenant_id=DEFAULT_TENANT_ID)
    out.seek(0)
    wb = load_workbook(out, read_only=True)
    return wb[wb.sheetnames[0]].cell(row=1, column=1).value


def test_classic_yearly_report_shows_the_saas_tenant_name(db, default_tenant, test_user, monkeypatch):
    monkeypatch.setattr(settings, "DEPLOYMENT_MODE", "saas")
    monkeypatch.setattr(settings, "PRACTICE_NAME", "Betreiber-Platzhalter")
    default_tenant.name = "Zahnarztpraxis Sonnenschein"
    default_tenant.slug = "zahnarztpraxis-sonnenschein"
    db.commit()

    assert _classic_a1(db) == "Zahnarztpraxis Sonnenschein"


def test_classic_yearly_report_on_prem_keeps_the_configured_name(db, default_tenant, test_user, monkeypatch):
    """On-prem unverändert: der Bootstrap-Mandant „Default" wird zum
    konfigurierten ``PRACTICE_NAME`` — derselbe Wert wie vor dem Fix."""
    monkeypatch.setattr(settings, "DEPLOYMENT_MODE", "onprem")
    monkeypatch.setattr(settings, "PRACTICE_NAME", DEMO_NAME)

    assert _classic_a1(db) == DEMO_NAME


def test_classic_yearly_report_neutralizes_a_formula_like_tenant_name(db, default_tenant, test_user, monkeypatch):
    """In SaaS ist der Name eine Eingabe aus dem öffentlichen Signup — er darf
    in der Tabelle nicht als Formel ausgeführt werden (wie jeder andere
    Nutzertext der Exporte)."""
    monkeypatch.setattr(settings, "DEPLOYMENT_MODE", "saas")
    default_tenant.name = '=HYPERLINK("http://example.invalid","Praxis")'
    default_tenant.slug = "formel-praxis"
    db.commit()

    assert _classic_a1(db) == '\'=HYPERLINK("http://example.invalid","Praxis")'


# ─── Vierte Fläche: die Anschrift im klassischen Jahresbericht (Zelle B1) ─
#
# Gleiche Lücke wie beim Namen, eine Zelle weiter: ``_create_employee_classic_sheet``
# schrieb ``settings.PRACTICE_ADDRESS`` direkt nach B1. On-prem ist das die
# Anschrift aus der Konfiguration (``[practice] address`` bzw. ``.env``) — in
# SaaS aber die Einstellung des Betreibers, gedruckt in den Jahresbericht JEDES
# Mandanten, neben dessen eigenem Praxisnamen in A1. Ein §16-Beleg der Praxis
# A trüge damit die Anschrift des Betreibers. ``practice_display_address`` ist
# das Gegenstück zu ``practice_display_name``: dieselbe Platzhalter-Regel,
# sonst die Anschrift aus den Abrechnungsdaten des Mandanten.

OPERATOR_ADDRESS = "Betreiberweg 9, 99999 Betreiberstadt"


def _classic_b1(db):
    from app.services.export_service import generate_yearly_report_classic
    from openpyxl import load_workbook
    from tests.conftest import DEFAULT_TENANT_ID

    out = generate_yearly_report_classic(db, 2026, tenant_id=DEFAULT_TENANT_ID)
    out.seek(0)
    wb = load_workbook(out, read_only=True)
    return wb[wb.sheetnames[0]].cell(row=1, column=2).value


def test_classic_yearly_report_shows_the_saas_tenant_address(db, default_tenant, test_user, monkeypatch):
    monkeypatch.setattr(settings, "DEPLOYMENT_MODE", "saas")
    monkeypatch.setattr(settings, "PRACTICE_ADDRESS", OPERATOR_ADDRESS)
    default_tenant.name = "Zahnarztpraxis Sonnenschein"
    default_tenant.slug = "zahnarztpraxis-sonnenschein"
    default_tenant.billing_address = {"street": "Hauptstr. 5", "zip": "12345", "city": "Musterstadt"}
    db.commit()

    assert _classic_b1(db) == "Hauptstr. 5, 12345 Musterstadt"


def test_classic_yearly_report_omits_the_operator_address_in_saas(db, default_tenant, test_user, monkeypatch):
    """Ohne hinterlegte Anschrift lieber eine leere Zelle als die des
    Betreibers — die stünde sonst unter dem Namen einer fremden Praxis."""
    monkeypatch.setattr(settings, "DEPLOYMENT_MODE", "saas")
    monkeypatch.setattr(settings, "PRACTICE_ADDRESS", OPERATOR_ADDRESS)
    default_tenant.name = "Zahnarztpraxis Sonnenschein"
    default_tenant.slug = "zahnarztpraxis-sonnenschein"
    db.commit()

    assert _classic_b1(db) is None


def test_classic_yearly_report_on_prem_keeps_the_configured_address(db, default_tenant, test_user, monkeypatch):
    """On-prem unverändert: die Anschrift aus der Konfiguration."""
    monkeypatch.setattr(settings, "DEPLOYMENT_MODE", "onprem")
    monkeypatch.setattr(settings, "PRACTICE_ADDRESS", "Musterstr. 1, 80000 München")

    assert _classic_b1(db) == "Musterstr. 1, 80000 München"


def test_classic_yearly_report_neutralizes_a_formula_like_address(db, default_tenant, test_user, monkeypatch):
    """Die Abrechnungsanschrift setzt die Mandanten-Admin per API — sie darf in
    der Tabelle nicht als Formel ausgeführt werden."""
    monkeypatch.setattr(settings, "DEPLOYMENT_MODE", "saas")
    default_tenant.name = "Formelpraxis"
    default_tenant.slug = "formelpraxis"
    default_tenant.billing_address = {"street": "=1+1", "zip": "12345", "city": "Musterstadt"}
    db.commit()

    assert _classic_b1(db) == "'=1+1, 12345 Musterstadt"


def _tenant_with_address(name="Praxis am Markt", slug="praxis-am-markt", **address):
    return Tenant(name=name, slug=slug, is_active=True, mode="multi",
                  billing_address=address or None)


def test_address_helper_uses_the_configuration_for_the_bootstrap_tenant(monkeypatch):
    from app.services.practice_name_service import practice_display_address

    monkeypatch.setattr(settings, "DEPLOYMENT_MODE", "onprem")
    monkeypatch.setattr(settings, "PRACTICE_ADDRESS", "  Musterstr. 1, 80000 München  ")

    assert practice_display_address(_tenant("Default")) == "Musterstr. 1, 80000 München"


def test_address_helper_omits_an_empty_configured_address(monkeypatch):
    from app.services.practice_name_service import practice_display_address

    monkeypatch.setattr(settings, "DEPLOYMENT_MODE", "onprem")
    monkeypatch.setattr(settings, "PRACTICE_ADDRESS", "   ")

    assert practice_display_address(_tenant("Default")) is None
    assert practice_display_address(None) is None


def test_address_helper_never_returns_the_operator_address_in_saas(monkeypatch):
    """Auch nicht für einen SaaS-Mandanten, der zufällig „Default"/``default``
    heißt — ``PRACTICE_ADDRESS`` gehört dort dem Betreiber."""
    from app.services.practice_name_service import practice_display_address

    monkeypatch.setattr(settings, "DEPLOYMENT_MODE", "saas")
    monkeypatch.setattr(settings, "PRACTICE_ADDRESS", OPERATOR_ADDRESS)

    assert practice_display_address(_tenant("Default")) is None
    assert practice_display_address(_tenant_with_address()) is None


def test_address_helper_formats_the_billing_address_on_one_line(monkeypatch):
    from app.services.practice_name_service import practice_display_address

    monkeypatch.setattr(settings, "DEPLOYMENT_MODE", "saas")

    assert practice_display_address(_tenant_with_address(
        street=" Hauptstr. 5 ", zip="12345", city="Musterstadt", country="DE",
    )) == "Hauptstr. 5, 12345 Musterstadt"
    # Fehlende oder explizit leere Teile (``None`` aus dem Schema) entfallen,
    # statt als „None" gedruckt zu werden.
    assert practice_display_address(_tenant_with_address(
        street=None, zip=None, city="Musterstadt",
    )) == "Musterstadt"
    assert practice_display_address(_tenant_with_address(street="", zip="", city="")) is None


def test_address_helper_after_anonymization_on_prem(monkeypatch):
    """Ein anonymisierter Default-Mandant (Name „[gelöscht]", Abrechnungsdaten
    geleert, Slug bleibt Bootstrap-Anker) bekommt nicht wieder die Anschrift
    aus der Konfiguration — wie beim Namen gilt der Ersatz nur dem Platzhalter."""
    from app.services.practice_name_service import practice_display_address

    monkeypatch.setattr(settings, "DEPLOYMENT_MODE", "onprem")
    monkeypatch.setattr(settings, "PRACTICE_ADDRESS", "Musterstr. 1, 80000 München")

    assert practice_display_address(_tenant("[gelöscht]")) is None


def test_avv_reads_the_billing_address_through_the_same_parts(monkeypatch):
    """Der AVV-Entwurf liest dieselben Felder (eine Quelle für „welche Teile hat
    die Abrechnungsanschrift"), mehrzeilig und mit Land."""
    from app.services.avv_generator import generate_avv_pdf

    monkeypatch.setattr(settings, "DEPLOYMENT_MODE", "saas")
    tenant = Tenant(name="Praxis am Markt", slug="praxis-am-markt", is_active=True, mode="multi",
                    country="DE",
                    billing_address={"street": "Hauptstr. 5", "zip": None, "city": "Musterstadt"})

    text = _pdf_text(BytesIO(generate_avv_pdf(tenant)))

    assert "Hauptstr. 5" in text
    assert "Musterstadt" in text
    assert "None" not in text
