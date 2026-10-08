"""#497 + #498: Tageszeilen der §16-Datei-Exporte (XLSX/PDF/ODS, Monat + Jahr).

#497 — Spalte „Differenz"/„Diff." an Krank-/Fortbildungstagen
---------------------------------------------------------------
SICK/TRAINING reduzieren das Soll nicht, sie werden dem Ist gutgeschrieben
(§ 3 EntgFG, ``credited_absences``). Die Summenzeilen der Exporte ziehen das
seit 1.18.0 aus ``get_monthly_actual``; die Tagesspalte „Differenz" rechnete
aber weiter ``Netto − Soll`` mit Netto = reiner Stempelzeit. Ein Kranktag stand
damit als −Tagessoll (rot) in der Datei, Σ „Differenz" widersprach dem
„Saldo Monat" desselben Blatts. „Netto (Std)" bleibt bewusst Stempelzeit
(§16-Nachweis der Anwesenheit) — nur die Saldo-Aussage je Tag zieht die
Gutschrift mit, aus derselben Quelle wie die Summenzeile
(``calculation_service.credited_absence_hours`` → ``credit_day_weight``).

#498 — geteilte Dienste
-----------------------
Mehrere Zeiteinträge an einem Tag wurden zu „erster Beginn – letztes Ende,
Pause Σ" zusammengezogen: 07:42–12:02 + 14:15–16:30 stand als 07:42–16:30,
Pause 0, Netto 6,58 — rechnerisch nicht nachvollziehbar und optisch ein
§4-ArbZG-Verstoß. Die Tabellen-Exporte hängen deshalb zwei Spalten AN
(„Unterbrechung (Min)", „Arbeitsblöcke"; Spalten 1–10 bleiben unverändert,
Kundenauswertungen lesen sie positionsweise), das PDF listet die Blöcke in
„Von"/„Bis" untereinander und weist die Unterbrechung in eigener Spalte aus.
"""
from datetime import date, datetime, time
from decimal import Decimal
from io import BytesIO

import pytest
from openpyxl import load_workbook

from app.models import Absence, AbsenceType, PublicHoliday, SystemSetting, TimeEntry
from app.services import calculation_service as cs
from app.services import export_service, ods_export_service
from tests.conftest import DEFAULT_TENANT_ID

YEAR, MONTH = 2026, 9
SPLIT_DAY = date(2026, 9, 3)      # Do — geteilter Dienst
WORK_DAY = date(2026, 9, 4)       # Fr — ein Block
SATURDAY = date(2026, 9, 5)       # Sa — Wochenendarbeit
EMPTY_DAY = date(2026, 9, 7)      # Mo — nichts erfasst
SICK_DAY = date(2026, 9, 21)      # Mo
TRAINING_DAY = date(2026, 9, 22)  # Di
HOLIDAY_SICK = date(2026, 9, 24)  # Do — Feiertag + Krank (Gewicht 0)

XLSX_HEADERS_1_TO_10 = [
    "Datum", "Wochentag", "Von", "Bis", "Pause (Min)",
    "Netto (Std)", "Soll (Std)", "Differenz", "Abwesenheit", "Bemerkung",
]


def _entry(db, user, d, start, end, break_min=0):
    e = TimeEntry(user_id=user.id, tenant_id=DEFAULT_TENANT_ID, date=d,
                  start_time=start, end_time=end, break_minutes=break_min)
    db.add(e)
    db.commit()
    return e


def _absence(db, user, d, typ, hours=8.0):
    a = Absence(user_id=user.id, tenant_id=DEFAULT_TENANT_ID, date=d, type=typ, hours=hours)
    db.add(a)
    db.commit()
    return a


@pytest.fixture
def month(db, test_user):
    """40 h/5 Tage (8 h Tagessoll). Ein Monat mit allem, was die Tagesspalten
    auseinanderlaufen liess."""
    _entry(db, test_user, SPLIT_DAY, time(7, 42), time(12, 2))
    _entry(db, test_user, SPLIT_DAY, time(14, 15), time(16, 30))
    _entry(db, test_user, WORK_DAY, time(8, 0), time(16, 30), 30)
    _entry(db, test_user, SATURDAY, time(9, 0), time(12, 0))
    _absence(db, test_user, SICK_DAY, AbsenceType.SICK)
    _absence(db, test_user, TRAINING_DAY, AbsenceType.TRAINING)
    db.add(PublicHoliday(tenant_id=DEFAULT_TENANT_ID, date=HOLIDAY_SICK,
                         name="Testfeiertag", year=YEAR))
    db.commit()
    _absence(db, test_user, HOLIDAY_SICK, AbsenceType.SICK)
    return test_user


def _month_saldo(db, user, year=YEAR, month_=MONTH) -> float:
    return float(cs.get_monthly_actual(db, user, year, month_)
                 - cs.get_monthly_target(db, user, year, month_))


def _year_saldo(db, user, year=YEAR) -> float:
    return sum(_month_saldo(db, user, year, m) for m in range(1, 13))


def _as_date(v):
    return v.date() if isinstance(v, datetime) else v


# --------------------------------------------------------------------------
# Lese-Helfer je Format
# --------------------------------------------------------------------------

def _xlsx_employee_sheet(bio: BytesIO, user, yearly=False):
    bio.seek(0)
    wb = load_workbook(bio)
    name = user.last_name[:20] if yearly else f"{user.last_name} {user.first_name}"[:31]
    return wb[name]


def _xlsx_day_rows(sheet):
    """{datum: [Zellwerte Spalte 1..12]} — nur echte Tageszeilen."""
    rows = {}
    for r in sheet.iter_rows(min_col=1, max_col=12, values_only=True):
        d = _as_date(r[0])
        if isinstance(d, date):
            rows[d] = list(r)
    return rows


def _ods_table(bio_or_doc, sheet_name):
    from odf.opendocument import load as odf_load
    from odf.table import Table
    if isinstance(bio_or_doc, BytesIO):
        bio_or_doc.seek(0)
        doc = odf_load(bio_or_doc)
    else:
        doc = bio_or_doc
    return [t for t in doc.spreadsheet.getElementsByType(Table)
            if t.getAttribute("name") == sheet_name][0]


def _ods_rows(table):
    """Liste von Zeilen; je Zelle (Text, value-Attribut)."""
    from odf.table import TableRow, TableCell
    from odf.teletype import extractText
    out = []
    for tr in table.getElementsByType(TableRow):
        out.append([(extractText(c), c.getAttribute("value"))
                    for c in tr.getElementsByType(TableCell)])
    return out


def _ods_day_rows(table):
    rows = {}
    for r in _ods_rows(table):
        if not r:
            continue
        try:
            d = datetime.strptime(r[0][0], "%d.%m.%Y").date()
        except ValueError:
            continue
        rows[d] = r
    return rows


def _ods_header(table):
    for r in _ods_rows(table):
        if r and r[0][0] == "Datum":
            return [c[0] for c in r]
    raise AssertionError("Kopfzeile nicht gefunden")


def _pdf_main_table(db, monkeypatch, year=YEAR, month_=MONTH):
    """Die Tabellen-Daten des PDF abgreifen, bevor reportlab sie rendert.

    ``generate_monthly_report_pdf`` importiert ``Table`` beim Aufruf aus
    ``reportlab.platypus`` — ein Unterklassen-Spion dort sieht die Zellen
    (Paragraph.text = Rohtext inkl. ``<br/>``)."""
    import reportlab.platypus as rp
    captured = []
    orig = rp.Table

    class _Spy(orig):
        def __init__(self, data, *a, **kw):
            captured.append(data)
            super().__init__(data, *a, **kw)

    monkeypatch.setattr(rp, "Table", _Spy)
    out = export_service.generate_monthly_report_pdf(db, year, month_, tenant_id=DEFAULT_TENANT_ID)
    out.seek(0)
    assert out.read(4) == b"%PDF"

    def _txt(c):
        return getattr(c, "text", c if isinstance(c, str) else "")

    main = next(t for t in captured if t and _txt(t[0][0]) == "Datum")
    header = [_txt(c) for c in main[0]]
    rows = {}
    for r in main[1:]:
        rows[datetime.strptime(_txt(r[0]), "%d.%m.%Y").date()] = [_txt(c) for c in r]
    return header, rows


# ==========================================================================
# #497 — Differenz an Gutschrift-Tagen
# ==========================================================================

class TestIssue497XlsxMonthly:
    def test_sick_and_training_rows_are_balance_neutral(self, db, month):
        sheet = _xlsx_employee_sheet(export_service.generate_monthly_report(
            db, YEAR, MONTH, tenant_id=DEFAULT_TENANT_ID), month)
        rows = _xlsx_day_rows(sheet)
        for d in (SICK_DAY, TRAINING_DAY):
            netto, soll, diff = rows[d][5], rows[d][6], rows[d][7]
            assert netto == pytest.approx(0.0), "Netto bleibt Stempelzeit (§16)"
            assert soll == pytest.approx(8.0), "Soll bleibt stehen (§3 EntgFG)"
            assert diff == pytest.approx(0.0), f"{d}: Gutschrift-Tag ist saldo-neutral, Diff={diff}"

    def test_holiday_with_sick_stays_zero(self, db, month):
        """Feiertag + Krank: Soll 0, Gutschrift-Gewicht 0 → Diff 0 (keine
        Phantom-Überstunden aus der Gutschrift)."""
        sheet = _xlsx_employee_sheet(export_service.generate_monthly_report(
            db, YEAR, MONTH, tenant_id=DEFAULT_TENANT_ID), month)
        assert _xlsx_day_rows(sheet)[HOLIDAY_SICK][7] == pytest.approx(0.0)

    def test_sum_of_daily_diff_equals_month_saldo(self, db, month):
        sheet = _xlsx_employee_sheet(export_service.generate_monthly_report(
            db, YEAR, MONTH, tenant_id=DEFAULT_TENANT_ID), month)
        total = sum(r[7] for r in _xlsx_day_rows(sheet).values())
        assert total == pytest.approx(_month_saldo(db, month), abs=0.005)


class TestIssue497XlsxYearly:
    def test_sick_row_and_sum_equal_year_saldo(self, db, month):
        sheet = _xlsx_employee_sheet(export_service.generate_yearly_report(
            db, YEAR, tenant_id=DEFAULT_TENANT_ID), month, yearly=True)
        rows = _xlsx_day_rows(sheet)
        assert rows[SICK_DAY][7] == pytest.approx(0.0)
        assert rows[TRAINING_DAY][7] == pytest.approx(0.0)
        total = sum(r[7] for r in rows.values())
        assert total == pytest.approx(_year_saldo(db, month), abs=0.01)


class TestIssue497Ods:
    def test_monthly_sick_row_and_sum_equal_saldo(self, db, month):
        table = _ods_table(ods_export_service.generate_monthly_report(
            db, YEAR, MONTH, tenant_id=DEFAULT_TENANT_ID), f"{month.last_name} {month.first_name}"[:31])
        rows = _ods_day_rows(table)
        assert float(rows[SICK_DAY][7][1]) == pytest.approx(0.0)
        assert float(rows[TRAINING_DAY][7][1]) == pytest.approx(0.0)
        assert float(rows[SICK_DAY][5][1]) == pytest.approx(0.0), "Netto bleibt Stempelzeit"
        total = sum(float(r[7][1]) for r in rows.values())
        assert total == pytest.approx(_month_saldo(db, month), abs=0.005)

    def test_monthly_weekend_work_shows_its_diff(self, db, month):
        """Parität zu XLSX/PDF: Samstagsarbeit ist +3 h im Saldo — die Spalte
        „Differenz" darf dort nicht 0 zeigen."""
        table = _ods_table(ods_export_service.generate_monthly_report(
            db, YEAR, MONTH, tenant_id=DEFAULT_TENANT_ID), f"{month.last_name} {month.first_name}"[:31])
        assert float(_ods_day_rows(table)[SATURDAY][7][1]) == pytest.approx(3.0)

    def test_yearly_sum_equals_year_saldo(self, db, month):
        table = _ods_table(ods_export_service.generate_yearly_report(
            db, YEAR, tenant_id=DEFAULT_TENANT_ID), f"{month.last_name} {month.first_name}"[:31])
        rows = _ods_day_rows(table)
        assert float(rows[SICK_DAY][7][1]) == pytest.approx(0.0)
        assert float(rows[SATURDAY][7][1]) == pytest.approx(3.0)
        total = sum(float(r[7][1]) for r in rows.values())
        assert total == pytest.approx(_year_saldo(db, month), abs=0.01)


class TestIssue497Pdf:
    def test_sick_row_and_sum_equal_saldo(self, db, month, monkeypatch):
        header, rows = _pdf_main_table(db, monkeypatch)
        col = header.index("Diff.")
        assert float(rows[SICK_DAY][col]) == pytest.approx(0.0)
        assert float(rows[TRAINING_DAY][col]) == pytest.approx(0.0)
        total = sum(float(r[col]) for r in rows.values())
        assert total == pytest.approx(_month_saldo(db, month), abs=0.005)


def test_half_special_day_sick_is_neutral_in_daily_diff(db, test_user):
    """24.12. als Halbtag: Soll 4, Gutschrift 8 × 0,5 = 4 → Diff 0 — derselbe
    #146-Faktor wie auf der Soll-Seite (credit_day_weight)."""
    for key, val in (("special_day_dec24_mode", "half_day"),
                     ("special_day_dec24_counts_as_vacation", "false")):
        db.add(SystemSetting(key=key, value=val, description=key, tenant_id=DEFAULT_TENANT_ID))
    db.commit()
    dec24 = date(2026, 12, 24)
    _absence(db, test_user, dec24, AbsenceType.SICK, hours=8.0)
    sheet = _xlsx_employee_sheet(export_service.generate_monthly_report(
        db, 2026, 12, tenant_id=DEFAULT_TENANT_ID), test_user)
    row = _xlsx_day_rows(sheet)[dec24]
    assert row[6] == pytest.approx(4.0)
    assert row[7] == pytest.approx(0.0)


LATE_FIRST_WORK_DAY = date(2026, 9, 16)   # Mi — Eintritt mitten im Monat
SICK_BEFORE_START = date(2026, 9, 8)      # Di — Krankmeldung VOR dem Eintritt


@pytest.fixture
def late_starter_sick(db, test_user):
    """Krankmeldung vor dem Eintritt (z. B. ``first_work_day`` nachträglich
    nach hinten verschoben). ``get_monthly_actual`` schließt die Gutschrift
    über das Beschäftigungsfenster aus (#195) — die Tagesspalte „Differenz"
    muss dasselbe tun, sonst stünden dort +8 h Phantom-Überstunden."""
    test_user.first_work_day = LATE_FIRST_WORK_DAY
    db.commit()
    _absence(db, test_user, SICK_BEFORE_START, AbsenceType.SICK, hours=8.0)
    return test_user


class TestIssue497CreditOutsideEmploymentWindow:
    """Review-Fund R3: der Guard ``credit = … if in_window else 0`` steht in
    FÜNF Tagesschleifen (XLSX Monat/Jahr, PDF, ODS Monat/Jahr). Ohne diesen
    Test bliebe die Suite grün, wenn er an einer Stelle verloren ginge.

    In XLSX und PDF trägt der Guard allein: ohne ihn stünde dort Diff +8. Die
    ODS-Blätter schreiben im Zweig „Außerhalb des Beschäftigungszeitraums" die
    Differenz zusätzlich fest auf 0 — dort pinnen die Tests das Verhalten, falls
    dieser Zweig einmal auf die berechnete Differenz umgestellt wird."""

    def test_xlsx_monthly(self, db, late_starter_sick):
        sheet = _xlsx_employee_sheet(export_service.generate_monthly_report(
            db, YEAR, MONTH, tenant_id=DEFAULT_TENANT_ID), late_starter_sick)
        rows = _xlsx_day_rows(sheet)
        assert rows[SICK_BEFORE_START][6] == pytest.approx(0.0), "kein Soll vor dem Eintritt"
        assert rows[SICK_BEFORE_START][7] == pytest.approx(0.0), "keine Gutschrift vor dem Eintritt"
        total = sum(r[7] for r in rows.values())
        assert total == pytest.approx(_month_saldo(db, late_starter_sick), abs=0.005)

    def test_xlsx_yearly(self, db, late_starter_sick):
        sheet = _xlsx_employee_sheet(export_service.generate_yearly_report(
            db, YEAR, tenant_id=DEFAULT_TENANT_ID), late_starter_sick, yearly=True)
        rows = _xlsx_day_rows(sheet)
        assert rows[SICK_BEFORE_START][7] == pytest.approx(0.0)
        total = sum(r[7] for r in rows.values())
        assert total == pytest.approx(_year_saldo(db, late_starter_sick), abs=0.01)

    def test_pdf(self, db, late_starter_sick, monkeypatch):
        header, rows = _pdf_main_table(db, monkeypatch)
        col = header.index("Diff.")
        assert float(rows[SICK_BEFORE_START][col]) == pytest.approx(0.0)
        total = sum(float(r[col]) for r in rows.values())
        assert total == pytest.approx(_month_saldo(db, late_starter_sick), abs=0.005)

    def test_ods_monthly(self, db, late_starter_sick):
        user = late_starter_sick
        table = _ods_table(ods_export_service.generate_monthly_report(
            db, YEAR, MONTH, tenant_id=DEFAULT_TENANT_ID), f"{user.last_name} {user.first_name}"[:31])
        rows = _ods_day_rows(table)
        assert float(rows[SICK_BEFORE_START][7][1]) == pytest.approx(0.0)
        total = sum(float(r[7][1]) for r in rows.values())
        assert total == pytest.approx(_month_saldo(db, user), abs=0.005)

    def test_ods_yearly(self, db, late_starter_sick):
        user = late_starter_sick
        table = _ods_table(ods_export_service.generate_yearly_report(
            db, YEAR, tenant_id=DEFAULT_TENANT_ID), f"{user.last_name} {user.first_name}"[:31])
        rows = _ods_day_rows(table)
        assert float(rows[SICK_BEFORE_START][7][1]) == pytest.approx(0.0)
        total = sum(float(r[7][1]) for r in rows.values())
        assert total == pytest.approx(_year_saldo(db, user), abs=0.01)


def test_credited_absence_hours_helper():
    """Die Tages-Gutschrift hat EINE Quelle: Σ hours × credit_day_weight, nur
    TRAINING/SICK."""
    class _A:
        def __init__(self, typ, hours):
            self.type, self.hours = typ, hours

    day = [_A(AbsenceType.SICK, 4.0), _A(AbsenceType.TRAINING, 2.5),
           _A(AbsenceType.VACATION, 4.0), _A(AbsenceType.OVERTIME, 8.0)]
    mon = date(2026, 9, 21)
    assert cs.credited_absence_hours(day, mon, set(), {}) == Decimal("6.5")
    assert cs.credited_absence_hours(day, mon, {mon}, {}) == Decimal("0")
    assert cs.credited_absence_hours(day, date(2026, 9, 19), set(), {}) == Decimal("0")  # Sa
    assert cs.credited_absence_hours([], mon, set(), {}) == Decimal("0")


# ==========================================================================
# #498 — geteilte Dienste
# ==========================================================================

class TestIssue498Helper:
    def test_split_shift(self):
        class _E:
            def __init__(self, s, e, b=0):
                self.start_time, self.end_time, self.break_minutes = s, e, b

        blocks, gap = export_service.day_work_blocks(
            [_E(time(7, 42), time(12, 2)), _E(time(14, 15), time(16, 30))])
        assert blocks == "07:42–12:02, 14:15–16:30"
        assert gap == 133

    def test_single_entry_has_no_block_list(self):
        class _E:
            def __init__(self, s, e, b=0):
                self.start_time, self.end_time, self.break_minutes = s, e, b

        assert export_service.day_work_blocks([_E(time(8, 0), time(16, 0), 30)]) == ("", 0)
        assert export_service.day_work_blocks([]) == ("", None)

    def test_overlapping_entries_have_no_gap(self):
        """Überlappende Einträge (A 08–17, B 12–14; Review 2026-06-23) erzeugen
        keine negative und keine erfundene Unterbrechung."""
        class _E:
            def __init__(self, s, e, b=0):
                self.start_time, self.end_time, self.break_minutes = s, e, b

        blocks, gap = export_service.day_work_blocks(
            [_E(time(8, 0), time(17, 0)), _E(time(12, 0), time(14, 0))])
        assert gap == 0
        assert blocks == "08:00–17:00, 12:00–14:00"

    def test_open_entry(self):
        class _E:
            def __init__(self, s, e, b=0):
                self.start_time, self.end_time, self.break_minutes = s, e, b

        blocks, gap = export_service.day_work_blocks(
            [_E(time(7, 42), time(12, 2)), _E(time(14, 15), None)])
        assert blocks == "07:42–12:02, 14:15–offen"
        assert gap == 133


def _minutes(hhmm: str) -> int:
    h, m = hhmm.split(":")
    return int(h) * 60 + int(m)


class TestIssue498Xlsx:
    @pytest.mark.parametrize("yearly", [False, True])
    def test_columns_are_appended_not_inserted(self, db, month, yearly):
        if yearly:
            bio = export_service.generate_yearly_report(db, YEAR, tenant_id=DEFAULT_TENANT_ID)
            header_row = 3
        else:
            bio = export_service.generate_monthly_report(db, YEAR, MONTH, tenant_id=DEFAULT_TENANT_ID)
            header_row = 4
        sheet = _xlsx_employee_sheet(bio, month, yearly=yearly)
        headers = [sheet.cell(row=header_row, column=c).value for c in range(1, 13)]
        assert headers[:10] == XLSX_HEADERS_1_TO_10, "Spalten 1–10 bleiben, wo sie sind"
        assert headers[10:] == ["Unterbrechung (Min)", "Arbeitsblöcke"]

    @pytest.mark.parametrize("yearly", [False, True])
    def test_split_shift_row(self, db, month, yearly):
        if yearly:
            bio = export_service.generate_yearly_report(db, YEAR, tenant_id=DEFAULT_TENANT_ID)
        else:
            bio = export_service.generate_monthly_report(db, YEAR, MONTH, tenant_id=DEFAULT_TENANT_ID)
        rows = _xlsx_day_rows(_xlsx_employee_sheet(bio, month, yearly=yearly))
        r = rows[SPLIT_DAY]
        von, bis, pause, netto = r[2], r[3], r[4], r[5]
        assert (von, bis, pause) == ("07:42", "16:30", 0), "Rahmen bleibt kompatibel"
        assert r[10] == 133
        assert r[11] == "07:42–12:02, 14:15–16:30"
        # Die Zeile geht jetzt auf: Bis − Von − Pause − Unterbrechung = Netto.
        assert (_minutes(bis) - _minutes(von) - pause - r[10]) / 60 == pytest.approx(netto, abs=0.01)
        # Ein-Block-Tag: Unterbrechung 0, keine Blockliste; leerer Tag: leer.
        assert rows[WORK_DAY][10] == 0
        assert rows[WORK_DAY][11] in (None, "")
        assert rows[EMPTY_DAY][10] is None


class TestIssue498Ods:
    @pytest.mark.parametrize("yearly", [False, True])
    def test_split_shift_row(self, db, month, yearly):
        if yearly:
            bio = ods_export_service.generate_yearly_report(db, YEAR, tenant_id=DEFAULT_TENANT_ID)
        else:
            bio = ods_export_service.generate_monthly_report(db, YEAR, MONTH, tenant_id=DEFAULT_TENANT_ID)
        table = _ods_table(bio, f"{month.last_name} {month.first_name}"[:31])
        header = _ods_header(table)
        assert header[:10] == XLSX_HEADERS_1_TO_10
        assert header[10:12] == ["Unterbrechung (Min)", "Arbeitsblöcke"]
        rows = _ods_day_rows(table)
        r = rows[SPLIT_DAY]
        assert (r[2][0], r[3][0], r[4][0]) == ("07:42", "16:30", "0")
        assert r[10][1] == "133"
        assert r[11][0] == "07:42–12:02, 14:15–16:30"
        assert rows[WORK_DAY][10][1] == "0"
        assert rows[WORK_DAY][11][0] == ""
        assert len(rows[EMPTY_DAY]) < 11 or rows[EMPTY_DAY][10][0] == ""


class TestIssue498Pdf:
    def test_blocks_are_listed_and_gap_shown(self, db, month, monkeypatch):
        header, rows = _pdf_main_table(db, monkeypatch)
        assert "Unterbr. (Min)" in header
        r = rows[SPLIT_DAY]
        von, bis = r[header.index("Von")], r[header.index("Bis")]
        assert von.split("<br/>") == ["07:42", "14:15"]
        assert bis.split("<br/>") == ["12:02", "16:30"]
        assert r[header.index("Unterbr. (Min)")] == "133"
        # Ein-Block-Tag unverändert
        w = rows[WORK_DAY]
        assert (w[header.index("Von")], w[header.index("Bis")]) == ("08:00", "16:30")
        assert w[header.index("Unterbr. (Min)")] == "0"
        assert rows[EMPTY_DAY][header.index("Unterbr. (Min)")] == ""
