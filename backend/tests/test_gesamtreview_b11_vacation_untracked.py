"""Gesamtreview 1.20 (B11, Produktentscheidung 2026-10-17): „ohne Stundenzählung"
entscheidet allein ``track_hours=False`` (Leitende, #191) — nicht ein heute auf 0
stehendes Tagessoll.

Bis hierher erkannten ``get_vacation_account``, ``absence_days`` und
``admin_change_requests._vacation_day_contribution`` den Fall am LIVE-Tagessoll
(``get_daily_target(user) <= 0``). Eine Person mit Stundenzählung in Elternzeit
(heute 0 Wochenstunden) zählte dadurch JEDE Urlaubszeile als ganzen Tag — auch an
Tagen ohne Soll — und jeden freien Sondertag mit Urlaubsabzug, während die
Buchungs- und Vorprüfpfade (``is_vacation_billable_day``) solche Tage längst nach
dem Snapshot überspringen. Jetzt zählen bei ``track_hours=True`` nur Tage mit Soll
laut datumsaufgelöstem Snapshot.
"""
from datetime import date
from decimal import Decimal

from app.models import Absence, AbsenceType, User, UserRole, WorkingHoursChange
from app.models.system_setting import SystemSetting
from app.routers.admin_change_requests import _vacation_day_contribution
from app.services import calculation_service
from tests.conftest import DEFAULT_TENANT_ID

MAR = date(2026, 3, 2)    # Montag, vor der Elternzeit (40 h)
AUG = date(2026, 8, 3)    # Montag, in der Elternzeit (0 h)
DEC24 = date(2026, 12, 24)  # Donnerstag


def _user(db, *, track_hours=True, weekly_hours=40.0, username="b11"):
    u = User(
        username=username, email=f"{username}@example.de", password_hash="h",
        first_name="A", last_name="B", role=UserRole.EMPLOYEE,
        weekly_hours=weekly_hours, vacation_days=30, work_days_per_week=5,
        track_hours=track_hours, is_active=True, tenant_id=DEFAULT_TENANT_ID,
    )
    db.add(u)
    db.commit()
    db.refresh(u)
    return u


def _parental_leave_user(db):
    """40 h bis Ende Juni, ab 01.07.2026 0 Wochenstunden (Elternzeit) — die
    User-Zeile spiegelt den heute gültigen Wert 0."""
    u = _user(db, weekly_hours=0.0)
    db.add(WorkingHoursChange(user_id=u.id, tenant_id=DEFAULT_TENANT_ID,
                              weekly_hours=40.0, effective_from=date(2026, 1, 1)))
    db.add(WorkingHoursChange(user_id=u.id, tenant_id=DEFAULT_TENANT_ID,
                              weekly_hours=0.0, effective_from=date(2026, 7, 1)))
    db.commit()
    return u


def _vac(db, u, d, hours, half_day=False):
    a = Absence(user_id=u.id, tenant_id=DEFAULT_TENANT_ID, date=d,
                type=AbsenceType.VACATION, hours=hours, half_day=half_day)
    db.add(a)
    db.commit()
    return a


def _free_dec24(db):
    for key, value in (("special_day_dec24_mode", "free"),
                       ("special_day_dec24_counts_as_vacation", "true")):
        db.add(SystemSetting(key=key, value=value, description=key, tenant_id=DEFAULT_TENANT_ID))
    db.commit()


def test_parental_leave_counts_only_days_with_target(db, default_tenant):
    """Urlaub im März (Soll 8 h) zählt, eine Urlaubszeile im August (Soll 0) und
    der freie 24.12. mit Urlaubsabzug nicht — bisher 3,0 Tage."""
    u = _parental_leave_user(db)
    _vac(db, u, MAR, 8.0)
    _vac(db, u, AUG, 0.0)
    _free_dec24(db)
    acc = calculation_service.get_vacation_account(db, u, 2026)
    assert acc["used_days"] == 1.0, acc
    assert acc["track_hours"] is True, acc


def test_parental_leave_absence_days(db, default_tenant):
    u = _parental_leave_user(db)
    rows = [_vac(db, u, MAR, 8.0), _vac(db, u, AUG, 0.0)]
    assert calculation_service.absence_days(db, u, rows) == Decimal("1")


def test_parental_leave_change_request_contribution(db, default_tenant):
    """Der Änderungsantrag prüft gegen das, was das Konto für die Zeile zählt."""
    u = _parental_leave_user(db)
    assert _vacation_day_contribution(db, u, AUG, False) == 0.0
    assert _vacation_day_contribution(db, u, MAR, False) == 1.0


def test_untracked_stays_day_based(db, default_tenant):
    """Kontrolle #191: ohne Stundenzählung zählt jede Zeile tagebasiert (voller
    Tag 1, Halbtag 0,5), der freie 24.12. mit Urlaubsabzug 1 — Stunden 0."""
    u = _user(db, track_hours=False, username="b11lead")
    rows = [_vac(db, u, MAR, 0.0), _vac(db, u, AUG, 0.0, half_day=True)]
    _free_dec24(db)
    acc = calculation_service.get_vacation_account(db, u, 2026)
    assert (acc["used_days"], acc["used_hours"], acc["track_hours"]) == (2.5, 0.0, False), acc
    assert calculation_service.absence_days(db, u, rows) == Decimal("1.5")
    assert _vacation_day_contribution(db, u, AUG, True) == 0.5


def test_normal_employee_unchanged(db, default_tenant):
    """Kontrolle: 40 h, zwei Urlaubstage + freier 24.12. → 3 Tage, 24 h."""
    u = _user(db, username="b11norm")
    rows = [_vac(db, u, MAR, 8.0), _vac(db, u, AUG, 8.0)]
    _free_dec24(db)
    acc = calculation_service.get_vacation_account(db, u, 2026)
    assert (acc["used_days"], acc["used_hours"], acc["track_hours"]) == (3.0, 24.0, True), acc
    assert calculation_service.absence_days(db, u, rows) == Decimal("2")
    assert _vacation_day_contribution(db, u, AUG, False) == 1.0


def test_parental_leave_hour_fields_not_negative(db, default_tenant):
    """Review-Fund zu B11: mit Stundenzählung und heute 0 h Tagessoll lassen sich
    die Urlaubstage nicht in Stunden umrechnen (budget_hours = Tage × 0). Die
    Stundenfelder bleiben deshalb wie vor B11 bei 0 — kein „Urlaub Rest (Std):
    −8,00" neben 29 offenen Tagen. Die Tage kommen aus dem Snapshot."""
    u = _parental_leave_user(db)
    _vac(db, u, MAR, 8.0)
    acc = calculation_service.get_vacation_account(db, u, 2026)
    assert (acc["used_days"], acc["remaining_days"]) == (1.0, 29.0), acc
    assert (acc["budget_hours"], acc["used_hours"], acc["remaining_hours"]) == (0.0, 0.0, 0.0), acc
    assert acc["track_hours"] is True, acc


def test_parental_leave_monthly_xlsx_no_negative_vacation_hours(db, default_tenant):
    """Datei-Export: das XLSX-Monatsblatt liest die Stundenfelder ungeprüft."""
    from io import BytesIO
    from openpyxl import load_workbook
    from app.services.export_service import generate_monthly_report

    u = _parental_leave_user(db)
    _vac(db, u, MAR, 8.0)
    bio = generate_monthly_report(db, 2026, 3)
    bio.seek(0)
    sheet = load_workbook(BytesIO(bio.read()))[f"{u.last_name} {u.first_name}"[:31]]
    values = {
        sheet.cell(row=r, column=1).value: sheet.cell(row=r, column=2).value
        for r in range(1, sheet.max_row + 1)
    }
    assert values["Urlaub Rest (Std):"] == 0.0, values
    assert values["Urlaub genommen (Std):"] == 0.0, values
