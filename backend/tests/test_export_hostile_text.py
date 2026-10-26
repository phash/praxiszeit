"""Sicherheits-Review 2026-10: Nutzer-Text darf keinen Tabellen-Export abschießen.

Fund 1: Steuerzeichen (\\x00-\\x08, \\x0b, \\x0c, \\x0e-\\x1f) in einer Notiz oder im
eigenen Namen (PUT /api/auth/profile) ließen openpyxl mit IllegalCharacterError
abbrechen → HTTP 500 für den XLSX-Export der GANZEN Praxis; das ODS bekam
ungültiges XML.

Fund 2: Blattnamen aus Nachnamen — ``\\ / * ? : [ ]`` lehnt openpyxl mit
ValueError ab („Müller/Schmidt" → jeder XLSX-Export 500), gleich benannte
Mitarbeitende kollidierten im ODS.
"""
import zipfile
from datetime import date, time
from xml.dom import minidom

import pytest
from openpyxl import load_workbook

from app.models import TimeEntry, User, UserRole
from app.services import auth_service, export_service, ods_export_service
from app.services.export_service import neutralize_spreadsheet_formula
from tests.conftest import DEFAULT_TENANT_ID

CTRL = "\x01\x00\x08\x0b\x0c\x1f"


def _mk_user(db, username, first, last):
    u = User(
        username=username,
        email=f"{username}@example.com",
        password_hash=auth_service.hash_password("testpassword123"),
        first_name=first,
        last_name=last,
        role=UserRole.EMPLOYEE,
        weekly_hours=40.0,
        vacation_days=30,
        work_days_per_week=5,
        is_active=True,
        tenant_id=DEFAULT_TENANT_ID,
    )
    db.add(u)
    db.commit()
    db.refresh(u)
    return u


def _entry(db, user, d, note):
    db.add(TimeEntry(user_id=user.id, tenant_id=DEFAULT_TENANT_ID, date=d,
                     start_time=time(8, 0), end_time=time(16, 0), break_minutes=30,
                     note=note))
    db.commit()


# --- Helfer ----------------------------------------------------------------

def test_neutralize_strips_illegal_control_chars():
    assert neutralize_spreadsheet_formula(f"a{CTRL}b") == "ab"


def test_neutralize_keeps_formula_semantics():
    assert neutralize_spreadsheet_formula("=1+1") == "'=1+1"
    assert neutralize_spreadsheet_formula("\t=x") == "'\t=x"
    # Steuerzeichen vor einer Formel: erst entfernen, dann neutralisieren
    assert neutralize_spreadsheet_formula("\x01=HYPERLINK()") == "'=HYPERLINK()"
    assert neutralize_spreadsheet_formula("Zeile1\nZeile2") == "Zeile1\nZeile2"
    assert neutralize_spreadsheet_formula(42) == 42
    assert neutralize_spreadsheet_formula(None) is None


def test_safe_sheet_title_rules():
    from app.services.export_service import safe_sheet_title
    assert safe_sheet_title("Müller/Schmidt Anna") == "Müller_Schmidt Anna"
    assert safe_sheet_title("a\\b*c?d:e[f]g") == "a_b_c_d_e_f_g"
    assert safe_sheet_title(f"X{CTRL}Y") == "XY"
    assert len(safe_sheet_title("A" * 50)) == 31
    assert safe_sheet_title("") == "Mitarbeiter"
    assert safe_sheet_title("  \x01 ") == "Mitarbeiter"
    assert safe_sheet_title("'Name'") == "Name"


def test_safe_sheet_title_deduplicates_case_insensitive():
    from app.services.export_service import safe_sheet_title
    existing = ["Meier Anna"]
    t = safe_sheet_title("meier anna", existing)
    assert t.lower() != "meier anna"
    assert len(t) <= 31
    long = "B" * 31
    t2 = safe_sheet_title(long, [long])
    assert t2 != long and len(t2) <= 31


# --- XLSX ------------------------------------------------------------------

@pytest.fixture
def hostile_users(db, default_tenant):
    a = _mk_user(db, "ma1", f"An{CTRL}na", "Müller/Schmidt")
    b = _mk_user(db, "ma2", "Anna", "Müller:Schmidt")  # gleicher sicherer Blattname
    _entry(db, a, date(2026, 1, 5), f"Notiz{CTRL}mit Steuerzeichen")
    _entry(db, b, date(2026, 1, 5), "=HYPERLINK(\"x\")")
    return a, b


def _xlsx_text(buf):
    wb = load_workbook(buf)
    texts = []
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            texts += [str(v) for v in row if v is not None]
    return wb, "\n".join(texts)


def test_xlsx_monthly_with_hostile_text(db, hostile_users):
    wb, text = _xlsx_text(export_service.generate_monthly_report(db, 2026, 1))
    assert "Notizmit Steuerzeichen" in text
    assert "'=HYPERLINK" in text
    names = wb.sheetnames
    assert len(names) == len(set(n.lower() for n in names)) >= 2
    assert all(not set(n) & set("\\/*?:[]") for n in names)


def test_xlsx_yearly_with_hostile_text(db, hostile_users):
    wb, _ = _xlsx_text(export_service.generate_yearly_report(db, 2026))
    assert len(wb.sheetnames) == len({n.lower() for n in wb.sheetnames})


def test_xlsx_classic_with_hostile_text(db, hostile_users):
    wb, _ = _xlsx_text(export_service.generate_yearly_report_classic(db, 2026))
    assert len(wb.sheetnames) == len({n.lower() for n in wb.sheetnames})


# --- ODS -------------------------------------------------------------------

def _ods_tables(buf):
    with zipfile.ZipFile(buf) as z:
        content = z.read("content.xml")
    dom = minidom.parseString(content)  # wirft bei ungültigem XML
    names = [t.getAttribute("table:name") for t in dom.getElementsByTagName("table:table")]
    for ch in CTRL:
        assert ch.encode() not in content
    return names


@pytest.mark.parametrize("gen", [
    lambda db: ods_export_service.generate_monthly_report(db, 2026, 1),
    lambda db: ods_export_service.generate_yearly_report(db, 2026),
    lambda db: ods_export_service.generate_yearly_report_classic(db, 2026),
])
def test_ods_with_hostile_text(db, hostile_users, gen):
    names = _ods_tables(gen(db))
    assert len(names) == len({n.lower() for n in names})
    assert all(not set(n) & set("\\/*?:[]") for n in names)
